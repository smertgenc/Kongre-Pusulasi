"""Read-only workbook audit and independent numerical verification."""
from pathlib import Path
import ast, hashlib, json, math, re
import numpy as np
import openpyxl
from engine import classify as classify_app

import argparse
_parser = argparse.ArgumentParser(description=__doc__)
_parser.add_argument('workbook', type=Path, help='Denetlenecek AHP Excel dosyası')
SOURCE = _parser.parse_args().workbook.expanduser().resolve()
OUT = Path(__file__).parent / 'analiz'
OUT.mkdir(exist_ok=True)
wf = openpyxl.load_workbook(SOURCE, data_only=False)
wc = openpyxl.load_workbook(SOURCE, data_only=True)
s, cached = wf['Sayfa1'], wc['Sayfa1']

def arithmetic(expr):
    node = ast.parse(expr, mode='eval')
    allowed = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant,
               ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd, ast.Load)
    if not all(isinstance(n, allowed) for n in ast.walk(node)):
        raise ValueError(expr)
    return eval(compile(node, '<numeric arithmetic>', 'eval'), {'__builtins__': {}})

memo = {}
def calc(addr):
    addr = addr.replace('$', '')
    if addr in memo:
        return memo[addr]
    v = s[addr].value
    if not isinstance(v, str) or not v.startswith('='):
        return v
    if v.startswith('=IF('):
        score = calc('L' + re.search(r'\d+', addr)[0])
        # Preserve the original workbook's text when verifying its cached formulas.
        result = '' if score == '' else ('Geçerli' if score >= 90 else 'Ortalama' if score >= 70 else 'Geçersiz')
    else:
        e = v[1:]
        e = re.sub(r'SUM\(([A-Z]+\d+):([A-Z]+\d+)\)',
                   lambda m: repr(sum(calc(c.coordinate) for row in s[m[1]+':'+m[2]] for c in row)), e)
        e = re.sub(r'\$?[A-Z]+\$?\d+', lambda m: repr(calc(m[0]) or 0), e)
        result = arithmetic(e)
    memo[addr] = result
    return result

def classify(x):
    return classify_app(x, {'average':70, 'valid':90})

def analyze(a):
    a = np.array(a, dtype=float)
    n = len(a)
    normal = a / a.sum(axis=0)
    w = normal.mean(axis=1)
    lam = float(np.mean(a @ w / w))
    ev, vectors = np.linalg.eig(a)
    ix = np.argmax(ev.real)
    ew = vectors[:, ix].real
    ew /= ew.sum()
    el = float(ev[ix].real)
    ri = {3: .58, 4: .90}[n]
    return dict(matrix=a.tolist(), normalized=normal.tolist(), weights=w.tolist(),
                weighted_sum=(a @ w).tolist(), ratios=(a @ w / w).tolist(),
                lambda_approx=lam, ci=(lam-n)/(n-1), cr=(lam-n)/(n-1)/ri,
                eigen_weights=ew.tolist(), lambda_eigen=el, cr_eigen=(el-n)/(n-1)/ri)

matrices = {}
for key, region in [('root','B3:D5'),('A','B22:D24'),('B','B41:E44'),('C','B62:D64')]:
    matrices[key] = analyze([[c.value for c in row] for row in s[region]])

weights = [calc(f'E{r}') for r in range(81,91)]
tests = []
for r in range(94,98):
    scores = [s.cell(r,c).value for c in range(2,12)]
    score = sum(x*w for x,w in zip(scores, weights))*100
    tests.append(dict(name=s[f'A{r}'].value, inputs=scores, expected=cached[f'L{r}'].value,
                      actual=score, delta=score-cached[f'L{r}'].value, status=classify(score)))

checks = []
for row in s:
    for c in row:
        if c.data_type == 'f':
            value = calc(c.coordinate)
            cv = cached[c.coordinate].value
            match = (value == cv) if isinstance(value,str) else math.isclose(value,cv,rel_tol=1e-12,abs_tol=1e-12)
            checks.append(dict(cell=c.coordinate, formula=c.value, recomputed=value, cached=cv, match=match))
assert all(x['match'] for x in checks)
assert math.isclose(sum(weights),1,abs_tol=1e-12)

mutated = np.ones((4,4)); mutated[0,3]=9; mutated[3,0]=1/9
norm = mutated/mutated.sum(axis=0)
mutation = dict(matrix=mutated.tolist(), legacy=norm[:,:3].mean(axis=1).tolist(),
                corrected=norm.mean(axis=1).tolist(), **{'diagnostic': analyze(mutated)})

scenarios = []
for name,x in [('all_one',[1]*10),('all_zero',[0]*10),('only_K1_K2',[1,1,0,0,0,0,0,0,0,0])]:
    z=float(np.dot(x,weights)*100)
    scenarios.append(dict(name=name,score=z,status=classify(z)))
from itertools import product
combinations = [(float(np.dot(x,weights)*100), x) for x in product([0,1],repeat=10)]
for t in [70,90]:
    for label, subset in [('below',[p for p in combinations if p[0]<t]),('above',[p for p in combinations if p[0]>=t])]:
        z,x=min(subset,key=lambda p:abs(p[0]-t))
        scenarios.append(dict(name=f'{t}_{label}',inputs=x,score=z,status=classify(z)))

whatif=[]
for code,idx,target in [('K3',2,.20),('K1',0,.20)]:
    changed=np.array(weights)*(1-target)/(1-weights[idx]); changed[idx]=target
    whatif.append(dict(code=code,target=target,weights=changed.tolist(),results=[dict(name=t['name'],score=float(np.dot(t['inputs'],changed)*100)) for t in tests]))

result=dict(source=SOURCE.name,sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    formula_count=len(checks),all_formulas_match=True,matrices=matrices,weights=weights,
    congresses=tests,mutation=mutation,scenarios=scenarios,whatif=whatif,
    dashboard_mean=float(np.mean([t['actual'] for t in tests])))
(OUT/'hesaplama_sonuclari.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))

def meaning(addr):
    col=re.match('[A-Z]+',addr)[0]; row=int(re.search(r'\d+',addr)[0])
    if row in [6,25,45,65]: return 'Sütun/ağırlık toplamını kontrol eder.'
    for start,end,wcol in [(3,5,'K'),(22,24,'K'),(41,44,'L'),(62,64,'K')]:
        if start<=row<=end:
            if col==wcol:
                return 'Normalize satır ortalaması → yerel ağırlık.' if start!=41 else 'HATA: Dört yerine üç normalize sütunun ortalaması → yerel ağırlık; K sütunu dışlanmış.'
            return 'Ham karşılaştırmayı kendi sütun toplamına böler → normalize matris.'
    for start,end in [(11,13),(30,32),(50,53),(70,72)]:
        if start<=row<=end:
            return {'B':'Ham matris satırı × ağırlık vektörü → (Aw)i.','C':'(Aw)i / wi → özdeğer yaklaşım bileşeni.','D':'Özdeğer yaklaşım bileşeni / n → ortalamaya katkı.'}[col]
    if row in [14,33,54,73]: return 'Bileşenlerin ortalaması → yaklaşık λmax; kesin özdeğer değildir.'
    if row in [15,34,55,74]: return '(λ−n)/(n−1) → CI.'
    if row in [16,35,56,75]: return 'CI / boyuta uygun RI → yaklaşık CR.'
    if 81<=row<=90:
        return {'C':'Üst grubun yerel ağırlığını getirir.','D':'Kriterin grup içi yerel ağırlığını getirir.','E':'Üst grup × yerel ağırlık → global ağırlık.'}[col]
    return 'Kriter skoru × global ağırlık toplamı ×100 → kongre puanı.' if col=='L' else '90 ve 70 eşikleriyle durum; puanın boş olup olmadığını kontrol eder, kriter eksiklerini kontrol etmez.'

lines=['# Excel hücre ve formül denetimi','',f'Kaynak SHA-256: `{result["sha256"]}`','',
       'Tek sayfa: Sayfa1. Dolu içerik 1–97. satırlar arasındadır; kullanılan alan biçimlendirme nedeniyle 110. satıra uzanır. Boş satırlar da tarandı. Kaynak dosya değiştirilmedi.',
       '',f'{len(checks)} formül bağımsız yeniden hesaplandı ve kayıtlı değerleriyle 1e-12 toleransta eşleşti. Bu işlem Excel uygulamasında yeniden hesaplama değildir. Matematiksel hata içeren formül de kendi kayıtlı sonucuyla eşleşebilir.','',
       '## Satır bazında bütün dolu hücreler','']
for row in s:
    cells=[c for c in row if c.value is not None]
    if not cells: continue
    lines.extend([f'### Satır {row[0].row}','','| Hücre | İçerik / formül | Kayıtlı sonuç | Matematiksel / yapısal anlam |','|---|---|---|---|'])
    for c in cells:
        val=str(c.value).replace('|','\\|').replace('\n',' ')
        desc=meaning(c.coordinate) if c.data_type=='f' else 'Etiket veya sabit girdi.'
        if c.coordinate in ['B17','B36','B57','B76','I81','I82','I83','I84']: desc='Sabit durum metni; CR değişince güncellenmez.'
        if c.coordinate in ['H81','H82','H83','H84']: desc='Elle girilmiş CR özeti; hesaplanan CR hücresine bağlı değil.'
        lines.append(f'| {c.coordinate} | `{val}` | {cached[c.coordinate].value if c.data_type=="f" else "—"} | {desc} |')
    lines.append('')
(OUT/'hucre_formul_denetimi.md').write_text('\n'.join(lines))
print(json.dumps(result,ensure_ascii=False,indent=2))
